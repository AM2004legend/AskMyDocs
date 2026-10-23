import pytest

from app.ingest.parsers import ParseError, parse_file
from tests.conftest import make_pdf


def test_txt_joins_wrapped_lines_into_paragraphs(tmp_path):
    path = tmp_path / "book.txt"
    path.write_text("CHAPTER I.\n\nIt is a truth universally\nacknowledged, that a man.\n\nSecond para.", "utf-8")
    parsed = parse_file(path, "txt")
    assert [b.text for b in parsed.blocks] == [
        "CHAPTER I.",
        "It is a truth universally acknowledged, that a man.",
        "Second para.",
    ]


def test_txt_strips_gutenberg_licence_and_finds_title(tmp_path):
    path = tmp_path / "pg.txt"
    path.write_text(
        "The Project Gutenberg eBook\n\nTitle: Tiny Book\n\n*** START OF THE PROJECT GUTENBERG EBOOK TINY ***\n\n"
        "The story.\n\n*** END OF THE PROJECT GUTENBERG EBOOK TINY ***\n\nSection 2. Licence stuff.",
        "utf-8",
    )
    parsed = parse_file(path, "txt")
    assert parsed.title == "Tiny Book"
    assert [b.text for b in parsed.blocks] == ["The story."]


def test_txt_falls_back_to_cp1252(tmp_path):
    path = tmp_path / "old.txt"
    path.write_bytes("caf\xe9 au lait".encode("cp1252"))
    assert parse_file(path, "txt").blocks[0].text == "café au lait"


def test_md_headings_have_levels(tmp_path):
    path = tmp_path / "notes.md"
    md = "# Title\n\n## Krebs cycle\nIt makes ATP.\nIn mitochondria.\n\n## Glycolysis\nIn cytoplasm."
    path.write_text(md, "utf-8")
    blocks = parse_file(path, "md").blocks
    assert [(b.text, b.level) for b in blocks] == [
        ("Title", 1),
        ("Krebs cycle", 2),
        ("It makes ATP. In mitochondria.", 0),
        ("Glycolysis", 2),
        ("In cytoplasm.", 0),
    ]


def test_docx_heading_styles_and_tables(tmp_path):
    import docx

    document = docx.Document()
    document.add_heading("Termination", level=1)
    document.add_paragraph("Either party may terminate with 30 days notice.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Notice period"
    table.rows[0].cells[1].text = "30 days"
    path = tmp_path / "contract.docx"
    document.save(path)

    blocks = parse_file(path, "docx").blocks
    assert (blocks[0].text, blocks[0].level) == ("Termination", 1)
    assert blocks[1].level == 0
    assert blocks[2].text == "Notice period | 30 days"


def test_epub_chapters_in_spine_order(tmp_path):
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("id1")
    book.set_title("Tiny Novel")
    chapters = []
    for n in (1, 2):
        ch = epub.EpubHtml(title=f"Chapter {n}", file_name=f"ch{n}.xhtml")
        ch.content = f"<h2>Chapter {n}</h2><p>Text of chapter {n}.</p>"
        book.add_item(ch)
        chapters.append(ch)
    book.toc = chapters
    book.spine = chapters
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    path = tmp_path / "novel.epub"
    epub.write_epub(str(path), book)

    parsed = parse_file(path, "epub")
    assert parsed.title == "Tiny Novel"
    texts = [(b.text, b.level) for b in parsed.blocks]
    assert ("Chapter 1", 2) in texts and ("Text of chapter 2.", 0) in texts
    assert texts.index(("Chapter 1", 2)) < texts.index(("Chapter 2", 2))


def test_pdf_keeps_page_numbers(tmp_path):
    path = tmp_path / "doc.pdf"
    path.write_bytes(make_pdf(["The first page talks about apples.", "The second page talks about pears."]))
    parsed = parse_file(path, "pdf")
    assert parsed.n_pages == 2
    assert parsed.ocr_pages == 0
    assert {b.page for b in parsed.blocks if "apples" in b.text} == {1}
    assert {b.page for b in parsed.blocks if "pears" in b.text} == {2}


def test_broken_pdf_gives_clear_error(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"this is not a pdf")
    with pytest.raises(ParseError):
        parse_file(path, "pdf")


def test_empty_file_is_rejected(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("   \n\n  ", "utf-8")
    with pytest.raises(ParseError, match="No readable text"):
        parse_file(path, "txt")


def test_unsupported_type(tmp_path):
    with pytest.raises(ParseError, match="Unsupported"):
        parse_file(tmp_path / "x.exe", "exe")
